"""팀 시트(xlsx로 받은 파일)나 CSV에서 업체명과 플레이스 링크를 읽는다."""

import csv
import unicodedata
from pathlib import Path

NAME_HEADERS = ("업체명", "name", "place_name")
# 앞에 있는 것부터 링크 열로 인정한다. '네이버지도에서 찾기'(검색 링크 열)는 제외된다.
LINK_HEADER_HINTS = ("플레이스 링크", "플레이스링크", "place_link", "링크", "link", "url")
REGION_HEADER_HINTS = ("장소", "지역", "region")
# 팀 시트의 장소 번호(혜화=1, 신촌=2, 강남=3, 종로=4) → 업체명 검색 때 쓰는 기준 좌표(위도,경도)
REGION_COORDS = {
    "1": "37.5822,127.0019",  # 혜화역
    "2": "37.5559,126.9368",  # 신촌역
    "3": "37.4979,127.0276",  # 강남역
    "4": "37.5704,126.9830",  # 종각역
}


def _norm(v):
    return "" if v is None else str(v).strip()


def _nfc(s):
    return unicodedata.normalize("NFC", _norm(s))


def _region_code(v):
    """2.0 → '2', 전각 '４' → '4'. 숫자가 아니면 정리한 글자를 그대로 둔다(예: '강서')."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = unicodedata.normalize("NFKC", _norm(v))
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def region_coords(code, default=None):
    return REGION_COORDS.get(_region_code(code), default)


def _pick_columns(header):
    lowered = [_nfc(h).lower() for h in header]
    name_idx = next((i for i, h in enumerate(lowered) if h in NAME_HEADERS), None)
    link_idx = None
    for hint in LINK_HEADER_HINTS:
        for i, h in enumerate(lowered):
            if i != name_idx and hint.lower() in h:
                link_idx = i
                break
        if link_idx is not None:
            break
    region_idx = next((i for i, h in enumerate(lowered)
                       if i not in (name_idx, link_idx) and any(hint in h for hint in REGION_HEADER_HINTS)), None)
    return name_idx, link_idx, region_idx


def _cell(row, idx):
    return row[idx] if idx is not None and idx < len(row) else None


def _rows_from_table(table):
    for start, header in enumerate(table[:5]):
        name_idx, link_idx, region_idx = _pick_columns(header)
        if name_idx is not None:
            break
    else:
        raise ValueError("머리글에서 '업체명' 열을 찾지 못했습니다.")
    out = []
    for row in table[start + 1:]:
        name = _norm(_cell(row, name_idx))
        if not name:
            continue
        out.append({"name": name, "link": _norm(_cell(row, link_idx)), "region": _region_code(_cell(row, region_idx))})
    return out


def read_places_with_sheet(path, sheet=None):
    """(업체 목록, 읽은 시트 이름)을 돌려준다. CSV면 시트 이름은 None.

    sheet를 주지 않으면 '업체명' 열이 있는 첫 시트를 읽는다. 팀 시트처럼 탭이 여러 개면 sheet로 탭 이름을 준다.
    """
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True, read_only=True)
        try:
            if sheet:
                wanted = _nfc(sheet)
                names = [n for n in wb.sheetnames if _nfc(n) == wanted]
                if not names:
                    raise ValueError(f"'{sheet}' 시트가 없습니다. 있는 시트: {', '.join(wb.sheetnames)}")
            else:
                names = wb.sheetnames
            last_error = None
            for name in names:
                table = [list(r) for r in wb[name].iter_rows(values_only=True)]
                try:
                    return _rows_from_table(table), name
                except ValueError as e:
                    last_error = e
            raise last_error or ValueError("읽을 시트가 없습니다.")
        finally:
            wb.close()
    with path.open(encoding="utf-8-sig", newline="") as f:
        table = [row for row in csv.reader(f)]
    return _rows_from_table(table), None


def read_places(path, sheet=None):
    return read_places_with_sheet(path, sheet)[0]
