"""수집 결과를 CSV(누적)와 엑셀(공유용)로 저장한다."""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

REVIEW_COLUMNS = [
    "collected_at", "place_name", "place_id", "review_id", "created", "created_raw", "visited",
    "rating", "body", "body_len", "photo_count", "origin", "visit_count", "keywords",
    "reviewer_review_count", "reviewer_photo_count", "reviewer_key",
    "event_mention", "service_mention",
]
KOREAN_HEADERS = {
    "collected_at": "수집 시각(KST)", "place_name": "업체명", "place_id": "플레이스 ID",
    "review_id": "리뷰 ID", "created": "작성일", "created_raw": "작성일(원문)", "visited": "방문일",
    "rating": "별점", "body": "본문", "body_len": "본문 길이", "photo_count": "사진·영상 수",
    "origin": "인증 수단", "visit_count": "방문 차수", "keywords": "선택 키워드",
    "reviewer_review_count": "리뷰어 리뷰 수", "reviewer_photo_count": "리뷰어 사진 수",
    "reviewer_key": "리뷰어 구분값",
    "event_mention": "이벤트 언급(확인 필요)", "service_mention": "서비스 언급(약한 신호)",
}
LOG_COLUMNS = [
    "collected_at", "place_name", "link", "place_id", "match", "place_title", "place_category", "road_address", "status",
    "category", "http_status", "reviews_found", "review_total", "coupon_count", "more_clicks",
    "final_url", "error",
]
LOG_HEADERS = {
    "collected_at": "수집 시각(KST)", "place_name": "업체명(시트)", "link": "입력 링크", "place_id": "플레이스 ID",
    "match": "찾은 방법", "place_title": "네이버 업체명", "place_category": "업종", "road_address": "도로명 주소", "status": "결과",
    "category": "주소 경로", "http_status": "HTTP", "reviews_found": "이번에 읽은 리뷰",
    "review_total": "방문자 리뷰 수(전체)", "coupon_count": "쿠폰 수", "more_clicks": "더보기 횟수",
    "final_url": "최종 주소", "error": "오류",
}
_WIDTHS = {"body": 60, "place_name": 22, "keywords": 30, "collected_at": 19, "created_raw": 12,
           "place_title": 22, "road_address": 34, "link": 30, "final_url": 30, "error": 30}


def load_keys(path):
    path = Path(path)
    if not path.exists() or path.stat().st_size == 0:
        return set()
    with path.open(encoding="utf-8-sig", newline="") as f:
        return {(r.get("place_id", ""), r.get("review_id", "")) for r in csv.DictReader(f)}


def append_rows(path, rows, columns):
    """CSV에 행을 덧붙인다. 새 파일일 때만 BOM과 머리글을 쓴다(엑셀에서 한글이 깨지지 않게)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    is_new = not path.exists() or path.stat().st_size == 0
    with path.open("a", encoding="utf-8-sig" if is_new else "utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        if is_new:
            writer.writeheader()
        for row in rows:
            writer.writerow({k: ("" if row.get(k) is None else row.get(k)) for k in columns})


def append_new_reviews(path, rows):
    """이미 저장된 (업체, 리뷰) 조합은 빼고 새 리뷰만 덧붙인 뒤 그 목록을 돌려준다."""
    keys = load_keys(path)
    new_rows = []
    for r in rows:
        key = (str(r.get("place_id", "")), str(r.get("review_id", "")))
        if key in keys:
            continue
        keys.add(key)
        new_rows.append(r)
    if new_rows:
        append_rows(path, new_rows, REVIEW_COLUMNS)
    return new_rows


def read_rows(path):
    path = Path(path)
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


NUMERIC_COLUMNS = {
    "body_len", "photo_count", "visit_count", "reviewer_review_count", "reviewer_photo_count",
    "event_mention", "service_mention", "reviews_found", "more_clicks", "http_status",
    "rating", "review_total", "coupon_count",
}


def _clean(value, column=None):
    if isinstance(value, str):
        if column in NUMERIC_COLUMNS and value.strip().lstrip("-").isdigit():
            return int(value.strip())  # CSV에서 읽은 숫자는 숫자로
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _fill_sheet(ws, rows, columns, headers):
    ws.append([headers.get(c, c) for c in columns])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, col in enumerate(columns, start=1):
            value = _clean(row.get(col), col)
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
                cell.data_type = "s"  # 본문이 '='로 시작해도 수식으로 바뀌지 않게
            if col == "body":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for c_idx, col in enumerate(columns, start=1):
        letter = ws.cell(row=1, column=c_idx).column_letter
        ws.column_dimensions[letter].width = _WIDTHS.get(col, 13)
    ws.freeze_panes = "C2"


def export_xlsx(rows, path, columns=REVIEW_COLUMNS, sheet_title="리뷰", places=None):
    """리뷰 시트를 만들고, places(업체별 결과)가 있으면 '업체' 시트를 덧붙인다."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    _fill_sheet(ws, rows, columns, KOREAN_HEADERS)
    if places is not None:
        _fill_sheet(wb.create_sheet("업체"), places, LOG_COLUMNS, LOG_HEADERS)
    wb.save(path)
    return path
