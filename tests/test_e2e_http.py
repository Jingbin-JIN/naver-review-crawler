"""기본(http) 엔진: 로컬 가짜 서버로 '업체명 검색 → 페이지 HTML 읽기 → 정리 → 저장' 전체 흐름을 확인한다.

브라우저가 필요 없다. 서버가 돌려주는 데이터는 실제 네이버 구조를 본뜬 가짜 데이터(tests/fixtures_real.py)다.
"""

import csv
import json
import os
import subprocess
import sys
import tempfile
import threading
import unittest
import urllib.parse
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook

from tests import fixtures_real as fr

ROOT = Path(__file__).resolve().parents[1]
CAPTCHA_ID = "9999999"


class _Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def _send(self, status, body, content_type):
        data = body.encode("utf-8")
        self.send_response(status)
        self.send_header("content-type", content_type)
        self.send_header("content-length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        if parsed.path == f"/restaurant/{fr.PLACE_ID}/review/visitor":
            self._send(200, fr.page_html(), "text/html; charset=utf-8")
        elif parsed.path == f"/restaurant/{CAPTCHA_ID}/review/visitor":
            self._send(200, fr.CAPTCHA_HTML, "text/html; charset=utf-8")
        elif parsed.path == "/p/api/search/instant-search":
            query = urllib.parse.parse_qs(parsed.query).get("query", [""])[0]
            payload = fr.search_payload() if "가짜식당" in query else {"place": [], "all": []}
            self._send(200, json.dumps(payload, ensure_ascii=False), "application/json; charset=utf-8")
        else:
            self._send(404, "<!doctype html><html><body>없는 페이지</body></html>", "text/html; charset=utf-8")


def _read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


class HttpEngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        base = f"http://127.0.0.1:{cls.server.server_address[1]}"
        cls.template = base + "/{category}/{place_id}/review/visitor?reviewSort=recent"
        cls.search_url = base + "/p/api/search/instant-search"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()

    def _cli(self, *args):
        cmd = [sys.executable, str(ROOT / "crawl.py"), *args, "--url-template", self.template,
               "--search-url", self.search_url, "--delay-min", "0", "--delay-max", "0"]
        env = dict(os.environ, PYTHONIOENCODING="utf-8")
        return subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=120, env=env)

    def _team_sheet(self, tmp):
        wb = Workbook()
        other = wb.active
        other.title = "팀원B"
        other.append(["업체명", "장소(혜화=1, 신촌=2, 강남=3, 종로=4)"])
        other.append(["다른 사람 가게", 2])
        ws = wb.create_sheet("팀원A")
        ws.append(["업체명", "리뷰 이벤트 조건", "장소(혜화=1, 신촌=2, 강남=3, 종로=4)", "리뷰이벤트여부 1일차 "])
        ws.append(["가짜식당 종각점", "포토", 4, 1])  # 링크 없음 → 업체명으로 찾음
        ws.append(["없는 가게", None, 4, 1])  # 검색 결과에 같은 이름 없음
        ws.append(["가짜식당", None, 4, 1])  # 비슷한 이름만 있음 → 고르지 않음
        path = tmp / "team.xlsx"
        wb.save(path)
        return path

    def test_run_team_sheet_with_name_search(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            out = tmp / "data"
            proc = self._cli("run", "--input", str(self._team_sheet(tmp)), "--sheet", "팀원A", "--out", str(out),
                             "--categories", "restaurant")
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            self.assertIn("읽은 시트: 팀원A (업체 3곳)", proc.stdout)
            rows = _read_csv(out / "reviews.csv")
            self.assertEqual(sorted(r["review_id"] for r in rows), ["aaa111", "bbb222"])
            self.assertTrue(all(r["place_id"] == fr.PLACE_ID and r["place_name"] == "가짜식당 종각점" for r in rows))
            logs = _read_csv(out / "run_log.csv")
            self.assertEqual([l["status"] for l in logs], ["ok", "name_not_found", "name_not_found"])
            first = logs[0]
            self.assertEqual((first["match"], first["place_title"], first["coupon_count"], first["review_total"]),
                             ("name", "가짜식당 종각점", "2", "1178"))
            book = load_workbook(next(out.glob("reviews_*.xlsx")))
            self.assertEqual(book.sheetnames, ["리뷰", "업체"])
            self.assertEqual(book["리뷰"].max_row, 3)
            self.assertEqual(book["업체"].max_row, 4)
            texts = [f.read_text(encoding="utf-8-sig") for f in out.rglob("*") if f.suffix in (".csv", ".json")]
            texts += [str(c.value) for name in book.sheetnames for row in book[name].iter_rows() for c in row]
            for nick in fr.NICKNAMES:
                self.assertFalse(any(nick in t for t in texts), nick)

    def test_run_links_and_captcha(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "data"
            link = f"https://m.place.naver.com/restaurant/{fr.PLACE_ID}/home"
            captcha = f"https://m.place.naver.com/restaurant/{CAPTCHA_ID}/home"
            proc = self._cli("run", "--place", link, "--place", captcha, "--out", str(out),
                             "--categories", "restaurant")
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            logs = _read_csv(out / "run_log.csv")
            self.assertEqual([l["status"] for l in logs], ["ok", "error"])
            self.assertEqual(logs[0]["match"], "link")
            self.assertIn("captcha", logs[1]["error"])
            self.assertEqual(len(_read_csv(out / "reviews.csv")), 2)

    def test_no_search_skips_rows_without_links(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            out = tmp / "data"
            proc = self._cli("run", "--input", str(self._team_sheet(tmp)), "--sheet", "팀원A", "--no-search",
                             "--out", str(out))
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            self.assertEqual([l["status"] for l in _read_csv(out / "run_log.csv")], ["no_link"] * 3)

    def test_probe_http(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "data"
            proc = self._cli("probe", "--place", f"https://m.place.naver.com/restaurant/{fr.PLACE_ID}/home",
                             "--out", str(out), "--categories", "restaurant")
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            self.assertIn("찾은 리뷰: 2건", proc.stdout)
            self.assertIn("쿠폰 수: 2", proc.stdout)
            self.assertIn("업체: 가짜식당 종각점", proc.stdout)
            summary = (out / "probe" / f"probe_{fr.PLACE_ID}_restaurant_summary.txt").read_text(encoding="utf-8")
            for nick in fr.NICKNAMES:
                self.assertNotIn(nick, summary)


if __name__ == "__main__":
    unittest.main()
