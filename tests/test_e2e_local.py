"""브라우저 엔진: 네이버 대신 로컬 가짜 페이지로 전체 흐름(페이지 열기 → 더보기 → 응답 수집 → 정리 → 저장)을 확인한다.

실제 네이버 구조와 같다는 보장은 아니다. 실제 구조는 PC에서 `crawl.py probe`로 확인한다.
"""

import csv
import json
import os
import subprocess
import sys
import tempfile
import threading
import unittest
from datetime import date
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook

from tests import fixtures

ROOT = Path(__file__).resolve().parents[1]
PLACE_ID = "1234567"


def _browser_available():
    try:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as p:
            p.chromium.launch().close()
        return True
    except Exception:
        return False


class _Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def _send(self, status, body, content_type):
        data = body.encode("utf-8")
        self.send_response(status)
        self.send_header("content-type", content_type)
        self.send_header("content-length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        path = self.path.split("?")[0]
        if path == f"/restaurant/{PLACE_ID}/review/visitor":
            self._send(200, fixtures.review_page_html(), "text/html; charset=utf-8")
        elif path == "/s/abc":  # 단축 링크 흉내: 플레이스 주소로 넘겨 준다
            self.send_response(302)
            self.send_header("location", f"/restaurant/{PLACE_ID}/home")
            self.send_header("content-length", "0")
            self.end_headers()
        else:
            self._send(404, "<!doctype html><html><body>없는 페이지</body></html>", "text/html; charset=utf-8")

    def do_POST(self):
        length = int(self.headers.get("content-length") or 0)
        self.rfile.read(length)
        if self.path.startswith("/graphql"):
            self._send(200, json.dumps(fixtures.graphql_payload(), ensure_ascii=False),
                       "application/json; charset=utf-8")
        else:
            self._send(404, "{}", "application/json")


def _read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


@unittest.skipUnless(_browser_available(), "Playwright Chromium이 없어서 건너뜀 (python -m playwright install chromium)")
class LocalEndToEndTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        port = cls.server.server_address[1]
        cls.base = f"http://127.0.0.1:{port}"
        cls.template = f"{cls.base}/{{category}}/{{place_id}}/review/visitor?reviewSort=recent"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()

    def _cli(self, *args):
        cmd = [sys.executable, str(ROOT / "crawl.py"), *args, "--engine", "browser", "--url-template", self.template,
               "--delay-min", "0", "--delay-max", "0", "--wait-ms", "500"]
        env = dict(os.environ, PYTHONIOENCODING="utf-8")  # 윈도우에서도 출력 인코딩을 맞춘다
        return subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=240, env=env)

    def test_fetch_place_collects_page_and_more_button(self):
        from naver_reviews import browser, extract

        with browser.open_browser() as context:
            cap = browser.fetch_place(context, PLACE_ID, category="restaurant", url_template=self.template,
                                      max_more=2, wait_ms=500)
        self.assertEqual(cap["error"], "")
        self.assertEqual(cap["http_status"], 200)
        self.assertIsInstance(cap["apollo_state"], dict)
        self.assertEqual(len(cap["graphql"]), 1)  # 목록 '펼쳐서 더보기'를 눌러 응답 1개
        self.assertGreaterEqual(cap["more_clicks"], 1)
        rows = extract.extract_all(cap["apollo_state"], cap["graphql"], today=date(2026, 10, 9))
        self.assertEqual(sorted(r["review_id"] for r in rows), ["r1", "r2", "r3"])

    def test_cli_run_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            wb = Workbook()
            other = wb.active
            other.title = "팀원B"  # 다른 사람 탭은 읽지 않아야 한다
            other.append(["업체명", "장소(혜화=1, 신촌=2, 강남=3, 종로=4)", "플레이스 링크"])
            other.append(["다른 사람 가게", 2, "https://m.place.naver.com/restaurant/7777777/home"])
            ws = wb.create_sheet("팀원A")
            ws.append(["업체명", "리뷰 이벤트 조건", "장소(혜화=1, 신촌=2, 강남=3, 종로=4)", "리뷰이벤트여부 1일차 ",
                       "플레이스 링크"])
            ws.append(["테스트 가게", None, 4, 1, f"https://m.place.naver.com/restaurant/{PLACE_ID}/home"])
            ws.append(["링크 없는 가게", "포토", 4, 1, None])
            ws.append(["이상한 링크 가게", None, 4, 1, "그냥 글자"])
            wb.save(tmp / "team.xlsx")
            out = tmp / "data"
            args = ("run", "--input", str(tmp / "team.xlsx"), "--sheet", "팀원A", "--no-search", "--out", str(out),
                    "--categories", "place", "restaurant")

            proc = self._cli(*args)
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            self.assertIn("읽은 시트: 팀원A (업체 3곳)", proc.stdout)
            rows = _read_csv(out / "reviews.csv")
            self.assertEqual(sorted(r["review_id"] for r in rows), ["r1", "r2", "r3"])
            self.assertTrue(all(r["place_name"] == "테스트 가게" and r["place_id"] == PLACE_ID for r in rows))
            logs = _read_csv(out / "run_log.csv")
            self.assertEqual([l["status"] for l in logs], ["ok", "no_link", "unresolved_link"])
            self.assertEqual(logs[0]["category"], "restaurant")  # place 주소는 404 → restaurant로 넘어감
            xlsx_files = list(out.glob("reviews_*.xlsx"))
            self.assertEqual(len(xlsx_files), 1)
            sheet = load_workbook(xlsx_files[0])["리뷰"]
            self.assertEqual(sheet.max_row, 4)

            texts = [f.read_text(encoding="utf-8-sig") for f in out.rglob("*") if f.suffix in (".csv", ".json")]
            texts += [str(c.value) for row in sheet.iter_rows() for c in row]
            for nick in fixtures.NICKNAMES:
                self.assertFalse(any(nick in t for t in texts), nick)
            self.assertTrue(any(out.glob("raw/*/1234567.json")))

            proc2 = self._cli(*args)  # 같은 리뷰를 다시 모아도 중복 저장하지 않는다
            self.assertEqual(proc2.returncode, 0, proc2.stdout + proc2.stderr)
            self.assertEqual(len(_read_csv(out / "reviews.csv")), 3)
            self.assertEqual(len(_read_csv(out / "run_log.csv")), 6)

    def test_cli_probe(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "data"
            proc = self._cli("probe", "--place", f"https://m.place.naver.com/restaurant/{PLACE_ID}/home",
                             "--out", str(out), "--categories", "restaurant")
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            self.assertIn("찾은 리뷰: 3건", proc.stdout)
            self.assertIn("GraphQL 작업 이름: getVisitorReviews", proc.stdout)
            probe_dir = out / "probe"
            saved = (probe_dir / f"probe_{PLACE_ID}_restaurant.json").read_text(encoding="utf-8")
            summary = (probe_dir / f"probe_{PLACE_ID}_restaurant_summary.txt").read_text(encoding="utf-8")
            self.assertIn("찾은 리뷰: 3건", summary)
            self.assertIn("'body': 'body'", summary)  # 필드 대응
            self.assertIn('"nickname": "str"', summary)  # 구조만 있고 값은 없음
            sheet = load_workbook(probe_dir / f"probe_{PLACE_ID}_restaurant_reviews.xlsx")["리뷰"]
            self.assertEqual(sheet.max_row, 4)
            cells = [str(c.value) for row in sheet.iter_rows() for c in row]
            for nick in fixtures.NICKNAMES:
                self.assertNotIn(nick, saved)
                self.assertNotIn(nick, summary)
                self.assertFalse(any(nick in c for c in cells), nick)

    def test_cli_run_with_place_links(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "data"
            share = f"[네이버지도]\n테스트 가게\n서울 어딘가\nhttps://m.place.naver.com/restaurant/{PLACE_ID}/home"
            proc = self._cli("run", "--place", share, "--place", "그냥 글자", "--out", str(out),
                             "--categories", "restaurant")
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            rows = _read_csv(out / "reviews.csv")
            self.assertEqual(sorted(r["review_id"] for r in rows), ["r1", "r2", "r3"])
            self.assertTrue(all(r["place_name"] == "테스트 가게" for r in rows))
            logs = _read_csv(out / "run_log.csv")
            self.assertEqual([l["status"] for l in logs], ["ok", "unresolved_link"])
            self.assertNotIn("\n", logs[0]["link"])  # 여러 줄 공유 문구도 로그에는 한 줄로

    def test_short_link_falls_back_to_browser(self):
        import crawl
        from naver_reviews import browser

        def broken(url, timeout):  # 맥 파이썬 인증서 문제 같은 상황
            raise OSError("certificate verify failed")

        with browser.open_browser() as context:
            follow = crawl._make_follow(context, primary=broken)
            final = follow(f"{self.base}/s/abc", 10)
        self.assertIn(f"/restaurant/{PLACE_ID}/home", final)


if __name__ == "__main__":
    unittest.main()
