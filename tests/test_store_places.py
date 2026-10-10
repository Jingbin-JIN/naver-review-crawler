import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from naver_reviews import places, store


def _row(pid, rid, body="본문"):
    return {"place_id": pid, "review_id": rid, "place_name": "가게", "body": body, "body_len": len(body),
            "photo_count": 1, "event_mention": 0, "service_mention": 0}


class StoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_append_dedup_and_single_bom(self):
        path = self.dir / "reviews.csv"
        first = store.append_new_reviews(path, [_row("1", "a"), _row("1", "b"), _row("1", "a")])
        second = store.append_new_reviews(path, [_row("1", "b"), _row("2", "b"), _row("1", "c")])
        self.assertEqual([r["review_id"] for r in first], ["a", "b"])
        self.assertEqual([(r["place_id"], r["review_id"]) for r in second], [("2", "b"), ("1", "c")])
        raw = path.read_bytes()
        self.assertEqual(raw.count(b"\xef\xbb\xbf"), 1)  # BOM은 맨 앞에 한 번만
        with path.open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual(len(rows), 4)
        self.assertEqual(list(rows[0].keys()), store.REVIEW_COLUMNS)

    def test_export_xlsx_safe_strings(self):
        rows = [_row("1", "a", "=1+1 수식처럼 보이는 본문"), _row("1", "b", "제어문자\x0b포함")]
        rows[0]["photo_count"] = "3"  # CSV에서 읽은 문자열 숫자
        path = store.export_xlsx(rows, self.dir / "out.xlsx")
        ws = load_workbook(path)["리뷰"]
        header = [c.value for c in ws[1]]
        self.assertEqual(header[0], "수집 시각(KST)")
        body_col = store.REVIEW_COLUMNS.index("body") + 1
        photo_col = store.REVIEW_COLUMNS.index("photo_count") + 1
        self.assertEqual(ws.cell(row=2, column=body_col).value, "=1+1 수식처럼 보이는 본문")
        self.assertEqual(ws.cell(row=2, column=body_col).data_type, "s")
        self.assertEqual(ws.cell(row=3, column=body_col).value, "제어문자포함")
        self.assertEqual(ws.cell(row=2, column=photo_col).value, 3)


class PlacesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_reads_checklist_layout(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "체크시트"
        ws.append(["번호", "업체명", "네이버지도에서 찾기", "플레이스 링크 (공유 링크 붙여넣기)", "리뷰 이벤트 조건"])
        ws.append([1, "가짜식당 종각점", '=HYPERLINK("https://map.naver.com/p/search/x","열기")',
                   "https://naver.me/abc", "포토"])
        ws.append([2, "가짜순대국 종각직영점", None, None, None])
        ws.append([None, None, None, None, None])
        wb.create_sheet("사용법").append(["설명"])
        path = self.dir / "check.xlsx"
        wb.save(path)
        rows = places.read_places(path)
        self.assertEqual(rows, [{"name": "가짜식당 종각점", "link": "https://naver.me/abc", "region": ""},
                                {"name": "가짜순대국 종각직영점", "link": "", "region": ""}])

    def test_reads_csv(self):
        path = self.dir / "places.csv"
        path.write_text("업체명,플레이스 링크\n가게A,1234567\n,\n", encoding="utf-8-sig")
        self.assertEqual(places.read_places(path), [{"name": "가게A", "link": "1234567", "region": ""}])

    def test_missing_name_column(self):
        path = self.dir / "bad.csv"
        path.write_text("이름,링크\nA,1\n", encoding="utf-8")
        with self.assertRaises(ValueError):
            places.read_places(path)

    def _team_workbook(self):
        """팀 시트와 같은 모양: 탭마다 열 순서가 다르고, 팀원A 탭 오른쪽 끝에 '플레이스 링크' 열을 추가한 상태."""
        wb = Workbook()
        ws = wb.active
        ws.title = "팀원B"
        ws.append(["업체명", "장소(혜화=1, 신촌=2, 강남=3, 종로=4)", "리뷰이벤트여부 1일차 ", "리뷰 이벤트 조건"])
        ws.append(["모꼬지해산물", 2, 1, None])
        ws = wb.create_sheet("팀원A")
        ws.append(["업체명", "리뷰 이벤트 조건", "장소(혜화=1, 신촌=2, 강남=3, 종로=4)", "리뷰이벤트여부 1일차 ",
                   "리뷰이벤트여부 2일차", "리뷰이벤트여부 3일차", "리뷰이벤트여부 4일차", "플레이스 링크"])
        ws.append(["가짜식당 종각점", None, 4, 1, 1, "１", None, "https://naver.me/abc"])
        ws.append(["가짜순대국 종각직영점", "포토", 4, 1, 1, "１", None, None])
        ws.append(["서현 양갈비 양꼬치", None, 4, 1, "오늘마감(한정수량)", "１", None, "1234567"])
        path = self.dir / "team.xlsx"
        wb.save(path)
        return path

    def test_reads_team_sheet_tab(self):
        path = self._team_workbook()
        rows, used = places.read_places_with_sheet(path, "팀원A")
        self.assertEqual(used, "팀원A")
        self.assertEqual(rows, [{"name": "가짜식당 종각점", "link": "https://naver.me/abc", "region": "4"},
                                {"name": "가짜순대국 종각직영점", "link": "", "region": "4"},
                                {"name": "서현 양갈비 양꼬치", "link": "1234567", "region": "4"}])

    def test_team_sheet_default_tab_and_unicode_tab_name(self):
        import unicodedata

        path = self._team_workbook()
        rows, used = places.read_places_with_sheet(path)  # 탭을 안 고르면 '업체명'이 있는 첫 탭
        self.assertEqual(used, "팀원B")
        self.assertEqual(rows, [{"name": "모꼬지해산물", "link": "", "region": "2"}])
        rows = places.read_places(path, unicodedata.normalize("NFD", "팀원A"))  # 맥에서 자모가 풀린 입력
        self.assertEqual(len(rows), 3)

    def test_unknown_tab_lists_existing_tabs(self):
        path = self._team_workbook()
        with self.assertRaises(ValueError) as cm:
            places.read_places(path, "없는탭")
        self.assertIn("팀원B", str(cm.exception))
        self.assertIn("팀원A", str(cm.exception))


if __name__ == "__main__":
    unittest.main()
